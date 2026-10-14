# Task: AI Dashboard & Reporting Automation Specialist

from datetime import datetime, timedelta
import numpy as np
import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd


def generate_telemetry_data(records: int = 60) -> pd.DataFrame:
    np.random.seed(42)
    base_date = datetime.now() - timedelta(days=records)
    dates = [base_date + timedelta(days=i) for i in range(records)]

    revenue = np.random.normal(loc=12000, scale=2500, size=records).clip(3000)
    ad_spend = revenue * np.random.uniform(0.15, 0.35, size=records)
    conversions = (revenue / np.random.uniform(40, 60, size=records)).astype(int)

    # Inietta anomalie intenzionali per l'analisi
    revenue[12] = 26500.0
    revenue[45] = 2100.0

    return pd.DataFrame(
        {
            "Data": [d.strftime("%Y-%m-%d") for d in dates],
            "Entrate": np.round(revenue, 2),
            "Spesa_ADV": np.round(ad_spend, 2),
            "Conversioni": conversions,
        }
    )


def compute_ai_insights(df: pd.DataFrame) -> dict:
    df["Data_dt"] = pd.to_datetime(df["Data"])
    df["ROAS"] = np.round(df["Entrate"] / df["Spesa_ADV"], 2)

    # Rilevamento anomalie (Z-Score > 2.0)
    rev_mean = df["Entrate"].mean()
    rev_std = df["Entrate"].std()
    df["Z_Score"] = np.round((df["Entrate"] - rev_mean) / rev_std, 2)
    df["Anomalia"] = df["Z_Score"].apply(
        lambda x: "Spike" if x > 2.0 else ("Dip" if x < -2.0 else "Normale")
    )

    tot_rev = float(df["Entrate"].sum())
    tot_adv = float(df["Spesa_ADV"].sum())
    avg_roas = float(tot_rev / tot_adv) if tot_adv != 0 else 0.0
    anomalies_count = int((df["Anomalia"] != "Normale").sum())

    # Generazione sintesi esecutiva
    max_rev_idx = df["Entrate"].idxmax()
    insights = [
        f"Volume Entrate Totali: € {tot_rev:,.2f} con ROAS medio pari a {avg_roas:.2f}x.",
        f"Rilevate {anomalies_count} anomalie statistiche significative (metodo Z-Score a 2-sigma).",
        f"Picco massimo di fatturato: € {df.loc[max_rev_idx, 'Entrate']:,.2f} registrato il {df.loc[max_rev_idx, 'Data']}.",
        "Azione suggerita: Isolare i pattern dei giorni di Spike per replicare l'efficienza ADV sui cluster stabili.",
    ]

    metrics = {
        "tot_rev": tot_rev,
        "tot_adv": tot_adv,
        "avg_roas": avg_roas,
        "anomalies": anomalies_count,
        "insights": insights,
    }
    return metrics


def apply_corporate_style(
    ws,
    header_fill: str = "1F4E79",
    accent_fill: str = "D9E1F2",
    alert_fill: str = "F8CBAD",
):
    font_family = "Segoe UI"
    header_font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
    data_font = Font(name=font_family, size=10)

    thin_border = Border(
        left=Side(style="thin", color="D3D3D3"),
        right=Side(style="thin", color="D3D3D3"),
        top=Side(style="thin", color="D3D3D3"),
        bottom=Side(style="thin", color="D3D3D3"),
    )

    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            if cell.row == 1:
                cell.fill = PatternFill(
                    start_color=header_fill,
                    end_color=header_fill,
                    fill_type="solid",
                )
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.font = data_font
                cell.border = thin_border
                if isinstance(cell.value, (int, float)):
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    if cell.column in [2, 3]:  # Entrate, Spesa_ADV
                        cell.number_format = "€ #,##0.00"
                    elif cell.column == 4:  # Conversioni
                        cell.number_format = "#,##0"
                    elif cell.column == 5:  # ROAS
                        cell.number_format = "0.00x"
                    elif cell.column == 6:  # Z_Score
                        cell.number_format = "0.00"
                else:
                    cell.alignment = Alignment(horizontal="center", vertical="center")

    # Formattazione condizionale per Spike e Dip
    spike_rule = CellIsRule(
        operator="equal",
        formula=['"Spike"'],
        fill=PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
        font=Font(color="006100", bold=True),
    )
    dip_rule = CellIsRule(
        operator="equal",
        formula=['"Dip"'],
        fill=PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
        font=Font(color="9C0006", bold=True),
    )

    status_col_letter =
