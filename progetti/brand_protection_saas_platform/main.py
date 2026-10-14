# Task: Brand Protection SaaS Platform

from dataclasses import dataclass
from datetime import datetime, timezone
from typing import List
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


@dataclass
class ThreatRecord:
    incident_id: str
    detected_at: datetime
    channel: str
    target_brand: str
    infringing_url: str
    threat_type: str
    risk_score: int
    takedown_status: str
    estimated_revenue_loss: float

    @property
    def severity(self) -> str:
        if self.risk_score >= 80:
            return "CRITICAL"
        elif self.risk_score >= 50:
            return "HIGH"
        elif self.risk_score >= 30:
            return "MEDIUM"
        return "LOW"


class BrandProtectionReportGenerator:
    # Palette UI
    COLOR_PRIMARY = "1E293B"       # Slate 800
    COLOR_SECONDARY = "475569"     # Slate 600
    COLOR_ACCENT = "0EA5E9"        # Sky 500
    COLOR_BG_LIGHT = "F8FAFC"      # Slate 50
    COLOR_TEXT_LIGHT = "FFFFFF"

    # Severity Colors (Fill / Text)
    SEVERITY_COLORS = {
        "CRITICAL": "FEE2E2",      # Red 100
        "HIGH": "FFEDD5",          # Orange 100
        "MEDIUM": "FEF9C3",        # Yellow 100
        "LOW": "DCFCE7",           # Green 100
    }
    SEVERITY_TEXT_COLORS = {
        "CRITICAL": "991B1B",      # Red 800
        "HIGH": "9A3412",          # Orange 800
        "MEDIUM": "854D0E",        # Yellow 800
        "LOW": "166534",           # Green 800
    }

    def __init__(self, brand_name: str):
        self.brand_name = brand_name
        self.workbook = openpyxl.Workbook()
        self.thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
