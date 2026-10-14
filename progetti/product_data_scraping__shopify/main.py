# Task: Product Data Scraping, Shopify & B2B Uploads, and Automated Competitor Price Monitoring

import logging
import os
import re
from datetime import datetime
from typing import Any, Dict, List, Optional

import pandas as pd
import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s")
logger = logging.getLogger(__name__)


class CompetitorScraper:
    """Scrapes competitor product data and prices."""

    def __init__(self, timeout: int = 10):
        self.timeout = timeout
        self.headers = {
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) "
                "Chrome/118.0.0.0 Safari/537.36"
            )
        }

    def fetch_price(self, url: str, selector: str) -> Optional[float]:
        if not url or pd.isna(url) or not selector or pd.isna(selector):
            return None
        try:
            response = requests.get(url, headers=self.headers, timeout=self.timeout)
            response.raise_for_status()
            soup = BeautifulSoup(response.content, "html.parser")
            element = soup.select_one(selector)
            if not element:
                logger.warning(f"Selector '{selector}' not found on {url}")
                return None

            raw_text = element.get_text(strip=True)
            cleaned = re.sub(r"[^\d,\.]", "", raw_text)
            if not cleaned:
                return None

            if "," in cleaned and "." in cleaned:
                if cleaned.rfind(",") > cleaned.rfind("."):
                    cleaned =
