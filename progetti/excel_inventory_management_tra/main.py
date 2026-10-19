# Task: Excel Inventory Management Tracker for E-commerce Business

# Requirements: pip install pandas openpyxl

import pandas as pd
import os
import logging
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

# Configure Logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s'
)

class InventoryManager:
    def __init__(self, inventory_file: str, sales_file: str, restock_file: str, output_dir: str):
        self.inventory_file = inventory_file
        self.sales_file = sales_file
        self.restock_file = restock_file
        self.output_dir = output_dir
        
        if not os.path.exists(self.output_dir):
            os.makedirs(self.output_dir)

    def generate_sample_data(self):
        """Generates sample Excel files if they do not exist for out-of-the-box testing."""
        logging.info("Checking for existing data files...")
        
        if not os.path.exists(self.inventory_file):
            inv_data = {
                'SKU': ['SKU001', 'SKU002', 'SKU003', 'SKU004', 'SKU005'],
                'Product Name': ['Wireless Mouse', 'Mechanical Keyboard', 'USB-C Hub', 'Laptop Stand', 'Webcam 1080p'],
                'Category': ['Electronics', 'Electronics', 'Accessories', 'Accessories', 'Electronics'],
                'Unit Price': [25.99, 89.99, 34.50, 45.00, 59.99],
                'Current Stock': [150, 40, 200, 15, 80],
                'Reorder Level': [30, 20, 50, 20, 25]
            }
            pd.DataFrame(inv_data).to_excel(self.inventory_file, index=False)
            logging.info(f"Created sample {self.inventory_file}")

        if not os.path.exists(self.sales_file):
            sales_data = {
                'Order ID': ['ORD101', 'ORD102', 'ORD103', 'ORD104'],
                'Date': ['2023-10-01', '2023-10-02', '2023-10-02', '2023-10-03'],
                'SKU': ['SKU001', 'SKU004', 'SKU002', 'SKU004'],
                'Quantity Sold': [5, 10, 2, 8]
            }
            pd.
